from __future__ import annotations

import base64
import csv
import io
import json
import mimetypes
import os
import re
import smtplib
import ssl
from datetime import date, datetime, time, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .access import create_password_reset, password_reset_url, write_local_recovery_file
from .cb1_models import CB1Message, CB1ProviderConnection, CB1SocialContent
from .cb1_services import decrypt, dumps, encrypt, loads, sha, utcnow as cb1_utcnow
from .cumulative_product_models import MessageDeliveryContext
from .client_admin_models import (
    CampaignExportPackage,
    ClientTrainingAssignment,
    ClientTrainingResource,
    EmailConnectionEvent,
    ForecastImportBatch,
    SocialGenerationMetadata,
    SolutionRequestPreference,
)
from .config import settings
from .db import get_db
from .models import (
    Account,
    Activity,
    Contact,
    ForecastMonth,
    ForecastVersion,
    Lead,
    Notification,
    Opportunity,
    PiqOpportunity,
    ServiceCatalog,
    SolutionRequest,
    Tenant,
    TenantService,
    User,
    WebsiteSite,
)
from .permissions import (
    is_global_admin,
    require_client_campaign_write,
    require_client_operational_write,
    require_client_website_write,
    require_tenant_access,
)
from .security import current_user, require_request_origin
from .services import audit
from .utils import model_dict

router = APIRouter(prefix="/api/v521", tags=["v5.2-client-administrator-correction"])

PLATFORMS = ("facebook", "instagram", "linkedin", "x")
PROVIDER_EXPORTS = {"generic", "mailchimp", "constant_contact", "hubspot", "brevo"}


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _client_admin_or_managed(user: User, tenant_id: str) -> None:
    if is_global_admin(user):
        require_client_operational_write(user, tenant_id)
        return
    if user.tenant_id != tenant_id or user.tenant_role != "CLIENT_ADMIN":
        raise HTTPException(status_code=403, detail="Client Administrator access required")


def _safe_url(value: str) -> str:
    value = value.strip()
    if not value:
        raise HTTPException(status_code=422, detail="Enter a website URL")
    if not urlparse(value).scheme:
        value = f"https://{value}"
    parsed = urlparse(value)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(status_code=422, detail="Enter a valid website address, such as www.example.com")
    return value.rstrip("/")


def _json_row(obj: Any, exclude: tuple[str, ...] = ()) -> dict[str, Any]:
    return model_dict(obj, exclude=set(exclude))


class WebsiteConnectionIn(BaseModel):
    mode: str = Field(pattern="^(external|managed)$")
    external_url: str = ""


class SocialGenerateIn(BaseModel):
    objective: str = Field(min_length=5, max_length=2000)
    audience: str = Field(min_length=2, max_length=1000)
    tone: str = Field(default="professional and approachable", max_length=160)
    call_to_action: str = Field(default="Learn more", max_length=300)
    keywords: list[str] = Field(default_factory=list)
    suggested_visual: str = Field(default="", max_length=1000)


class CampaignExportIn(BaseModel):
    name: str = Field(min_length=3, max_length=220)
    provider_format: str = "generic"
    include_leads: bool = True
    include_contacts: bool = True
    include_piq: bool = False
    lead_statuses: list[str] = Field(default_factory=list)
    sources: list[str] = Field(default_factory=list)
    subject: str = Field(default="", max_length=300)
    preview_text: str = Field(default="", max_length=500)
    email_body: str = ""
    call_to_action: str = Field(default="", max_length=300)


class OneToOneEmailIn(BaseModel):
    connection_id: str
    recipient_email: str
    recipient_name: str = ""
    subject: str = Field(min_length=1, max_length=300)
    body: str = Field(min_length=1)
    account_id: str | None = None
    opportunity_id: str | None = None
    lead_id: str | None = None
    contact_id: str | None = None


class ForecastCommitIn(BaseModel):
    batch_id: str


class TrainingCompleteIn(BaseModel):
    status: str = Field(pattern="^(not_started|in_progress|complete)$")
    progress_pct: int = Field(default=100, ge=0, le=100)


class PasswordResetSendIn(BaseModel):
    user_id: str


class SolutionRequestIn(BaseModel):
    service_code: str
    note: str = Field(default="", max_length=5000)
    preferred_contact_method: str = Field(default="Email", max_length=40)
    contact_date: date
    contact_time: str = Field(pattern=r"^([01]\d|2[0-3]):[0-5]\d$")
    timezone: str = Field(min_length=2, max_length=100)


@router.get("/tenants/{tenant_id}/environment")
def environment_status(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    return {
        "release": "5.4.1.2-interaction-regression-correction-po1",
        "demo_mode": settings.install_profile == "demo" or settings.auto_seed,
        "ai_provider": os.getenv("RMR_AI_PROVIDER", "demonstration").lower(),
        "one_to_one_email": "client_connected_mailbox",
        "bulk_email_delivery": False,
        "campaign_export_formats": sorted(PROVIDER_EXPORTS),
        "external_pi_q_dependency": {
            "discovery_provider": os.getenv("RMR_PIQ_DISCOVERY_PROVIDER", "demonstration"),
            "adaptive_research_provider": os.getenv("RMR_PIQ_RESEARCH_PROVIDER", "demonstration"),
        },
    }


@router.get("/tenants/{tenant_id}/website-connection")
def website_connection(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    tenant = db.get(Tenant, tenant_id)
    site = db.scalar(select(WebsiteSite).where(WebsiteSite.tenant_id == tenant_id))
    if not tenant or not site:
        raise HTTPException(status_code=404, detail="Website configuration not found")
    return {
        "mode": "external" if site.mode == "external" else "managed",
        "external_url": site.external_url or tenant.website_url,
        "managed_url": f"{settings.base_url}/sites/{site.slug}",
        "status": site.status,
        "strategy": "Clients may connect an existing/custom website or use the preserved RMR managed-site capability.",
    }


@router.patch("/tenants/{tenant_id}/website-connection")
def update_website_connection(tenant_id: str, payload: WebsiteConnectionIn, request: Request,
                              user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    require_client_website_write(user, tenant_id)
    tenant = db.get(Tenant, tenant_id)
    site = db.scalar(select(WebsiteSite).where(WebsiteSite.tenant_id == tenant_id))
    if not tenant or not site:
        raise HTTPException(status_code=404, detail="Website configuration not found")
    if payload.mode == "external":
        external_url = _safe_url(payload.external_url)
        site.mode = "external"
        site.external_url = external_url
        tenant.website_mode = "external"
        tenant.website_url = external_url
    else:
        site.mode = "managed"
        tenant.website_mode = "managed"
        tenant.managed_site_slug = site.slug
    audit(db, user, "website.connection.updated", tenant_id=tenant_id, entity_type="website_site", entity_id=site.id,
          data={"mode": site.mode, "external_url": site.external_url})
    db.commit()
    return website_connection(tenant_id, user, db)


@router.post("/tenants/{tenant_id}/website-connection/test")
def test_website_connection(tenant_id: str, payload: WebsiteConnectionIn, request: Request,
                            user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    require_client_website_write(user, tenant_id)
    if payload.mode == "managed":
        site = db.scalar(select(WebsiteSite).where(WebsiteSite.tenant_id == tenant_id))
        if not site:
            raise HTTPException(status_code=404, detail="Managed site is not configured")
        return {"ok": True, "url": f"{settings.base_url}/sites/{site.slug}", "detail": "Managed website route is configured."}
    url = _safe_url(payload.external_url)
    live_test = os.getenv("RMR_EXTERNAL_SITE_LIVE_TEST", "false").lower() == "true"
    if not live_test:
        return {"ok": True, "url": url, "detail": "URL format is valid. Live reachability testing is disabled in this Product Owner environment."}
    try:
        response = httpx.get(url, follow_redirects=True, timeout=10)
        return {"ok": response.status_code < 500, "url": str(response.url), "status_code": response.status_code,
                "detail": f"External website returned HTTP {response.status_code}."}
    except Exception as exc:
        raise HTTPException(status_code=409, detail=f"External website could not be reached: {exc}")


def _demo_social(payload: SocialGenerateIn) -> dict[str, dict[str, Any]]:
    """High-quality deterministic Product Owner content.

    This is deliberately labelled demonstration output by the API. It proves
    the workflow without pretending that an external AI provider was called.
    """
    objective = payload.objective.strip().rstrip(".")
    audience = payload.audience.strip()
    tone = payload.tone.strip() or "professional and approachable"
    cta = payload.call_to_action.strip() or "Learn more"
    keywords = [k.strip().lstrip("#") for k in payload.keywords if k.strip()]
    if not keywords:
        stop = {"about", "their", "there", "these", "those", "would", "could", "should", "with", "from", "into", "your"}
        keywords = [
            w.lower() for w in re.findall(r"[A-Za-z][A-Za-z0-9]+", f"{objective} {audience}")
            if len(w) > 4 and w.lower() not in stop
        ][:7]
    tags = [f"#{re.sub(r'[^A-Za-z0-9]', '', k.title())}" for k in keywords if re.sub(r'[^A-Za-z0-9]', '', k)]
    visual = payload.suggested_visual.strip() or (
        f"An authentic, brand-consistent image showing the real-world outcome of {objective.lower()} for {audience}."
    )
    promise = objective[0].upper() + objective[1:] if objective else "A better next step"
    facebook_hook = f"A move this important deserves more than a generic checklist."
    facebook_post = "\n\n".join([
        facebook_hook,
        f"{promise}. If you are part of {audience}, the right plan can make the process feel far more manageable—from the first questions through the final decision.",
        "Start with a conversation about your goals, timing, and the details that matter most to you.",
        cta,
    ])
    instagram_hook = f"Your next chapter should feel exciting—not overwhelming. ✨"
    instagram_post = "\n\n".join([
        instagram_hook,
        f"{promise} for {audience}.",
        "A clear plan. Local guidance. Fewer surprises. More confidence in every next step.",
        f"👉 {cta}",
    ])
    linkedin_hook = f"The strongest outcomes usually begin with a better process."
    linkedin_post = "\n\n".join([
        linkedin_hook,
        f"For {audience}, {objective.lower()} is not only a transaction—it is a sequence of decisions that benefits from relevant market context, clear expectations, and consistent follow-through.",
        "A thoughtful process helps people move from uncertainty to an informed next step while keeping the experience personal.",
        cta,
    ])
    x_hook = f"A clearer plan makes the next move easier."
    x_post = f"{x_hook} {promise} for {audience}. {cta}"
    return {
        "facebook": {
            "title": facebook_hook, "post": facebook_post,
            "hashtags": " ".join((tags + ["#TrustedGuidance", "#NextStep"])[:7]),
            "hook": facebook_hook, "cta": cta, "keywords": keywords, "suggested_visual": visual,
        },
        "instagram": {
            "title": instagram_hook, "post": instagram_post,
            "hashtags": " ".join((tags + ["#NewChapter", "#LocalExpert", "#MoveWithConfidence", "#RealEstateGoals"])[:12]),
            "hook": instagram_hook, "cta": cta, "keywords": keywords, "suggested_visual": visual,
        },
        "linkedin": {
            "title": linkedin_hook, "post": linkedin_post,
            "hashtags": " ".join((tags + ["#ClientExperience", "#ProfessionalServices"])[:5]),
            "hook": linkedin_hook, "cta": cta, "keywords": keywords, "suggested_visual": visual,
        },
        "x": {
            "title": x_hook, "post": x_post[:270],
            "hashtags": " ".join((tags + ["#NextStep"])[:3]),
            "hook": x_hook, "cta": cta, "keywords": keywords, "suggested_visual": visual,
        },
    }


def _extract_json_object(text_value: str) -> dict[str, Any]:
    text_value = text_value.strip()
    if text_value.startswith("```"):
        text_value = re.sub(r"^```(?:json)?\s*|\s*```$", "", text_value, flags=re.I | re.S)
    start, end = text_value.find("{"), text_value.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("AI provider did not return a JSON object")
    return json.loads(text_value[start:end + 1])


def _real_social(payload: SocialGenerateIn) -> dict[str, dict[str, Any]]:
    provider = os.getenv("RMR_AI_PROVIDER", "demonstration").lower()
    api_key = os.getenv("RMR_AI_API_KEY") or os.getenv("OPENAI_API_KEY")
    if provider in {"demonstration", "mock", "local"}:
        return _demo_social(payload)
    if provider not in {"openai", "openai_compatible"}:
        raise HTTPException(status_code=503, detail=f"Unsupported AI provider: {provider}")
    if not api_key:
        raise HTTPException(status_code=503, detail="AI provider credentials are not configured")
    base_url = os.getenv("RMR_AI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("RMR_AI_MODEL", "gpt-5-mini")
    prompt = {
        "objective": payload.objective,
        "audience": payload.audience,
        "tone": payload.tone,
        "call_to_action": payload.call_to_action,
        "keywords": payload.keywords,
        "suggested_visual_context": payload.suggested_visual,
        "requirements": {
            "platforms": list(PLATFORMS),
            "for_each_platform": ["title", "post", "hashtags", "hook", "cta", "keywords", "suggested_visual"],
            "rules": [
                "Write finished, useful, genuinely differentiated platform-specific copy.",
                "Do not repeat the same paragraph with superficial platform changes.",
                "Do not invent factual claims, statistics, credentials, guarantees, or testimonials.",
                "X post must fit normal X length constraints.",
                "Return JSON only with a top-level platforms object.",
            ],
        },
    }
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    if provider == "openai":
        response = httpx.post(
            f"{base_url}/responses",
            headers=headers,
            json={
                "model": model,
                "store": False,
                "instructions": "Create human-reviewed, platform-specific social marketing content. Return only valid JSON.",
                "input": json.dumps(prompt),
            },
            timeout=90,
        )
        response.raise_for_status()
        data = response.json()
        content = data.get("output_text") or ""
        if not content:
            chunks = []
            for output in data.get("output", []):
                for part in output.get("content", []):
                    if part.get("type") in {"output_text", "text"} and part.get("text"):
                        chunks.append(part["text"])
            content = "\n".join(chunks)
    else:
        response = httpx.post(
            f"{base_url}/chat/completions",
            headers=headers,
            json={
                "model": model,
                "messages": [
                    {"role": "system", "content": "Create human-reviewed, platform-specific social marketing content. Return only valid JSON."},
                    {"role": "user", "content": json.dumps(prompt)},
                ],
                "response_format": {"type": "json_object"},
            },
            timeout=90,
        )
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"]
    data = _extract_json_object(content)
    result = data.get("platforms", data)
    missing = [platform for platform in PLATFORMS if platform not in result]
    if missing:
        raise HTTPException(status_code=502, detail=f"AI response omitted platforms: {', '.join(missing)}")
    for platform in PLATFORMS:
        if not str(result[platform].get("post") or "").strip():
            raise HTTPException(status_code=502, detail=f"AI response returned empty {platform} copy")
    return {platform: result[platform] for platform in PLATFORMS}


@router.post("/tenants/{tenant_id}/social/generate")
def generate_social(tenant_id: str, payload: SocialGenerateIn, request: Request,
                    user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    require_client_campaign_write(user, tenant_id)
    generated = _real_social(payload)
    provider = os.getenv("RMR_AI_PROVIDER", "demonstration").lower()
    items = []
    for platform in PLATFORMS:
        item = generated[platform]
        row = CB1SocialContent(
            tenant_id=tenant_id,
            platform=platform.upper(),
            title=str(item.get("title") or item.get("hook") or payload.objective)[:240],
            post_text=str(item.get("post") or ""),
            hashtags=str(item.get("hashtags") or ""),
            status="DRAFT",
            created_by=user.id,
        )
        db.add(row)
        db.flush()
        db.add(SocialGenerationMetadata(
            social_content_id=row.id,
            tenant_id=tenant_id,
            provider=provider,
            objective=payload.objective,
            audience=payload.audience,
            tone=payload.tone,
            hook=str(item.get("hook") or ""),
            call_to_action=str(item.get("cta") or payload.call_to_action),
            keywords_json=[str(x) for x in item.get("keywords", payload.keywords)],
            suggested_visual=str(item.get("suggested_visual") or payload.suggested_visual),
            metadata_json={"platform": platform},
        ))
        items.append({**_json_row(row), "generation": {
            "provider": provider,
            "hook": str(item.get("hook") or ""),
            "cta": str(item.get("cta") or payload.call_to_action),
            "keywords": item.get("keywords", payload.keywords),
            "suggested_visual": str(item.get("suggested_visual") or payload.suggested_visual),
        }})
    audit(db, user, "social.ai.generated", tenant_id=tenant_id, entity_type="social_content", data={"provider": provider, "objective": payload.objective})
    db.commit()
    return {"items": items, "provider": provider, "demonstration": provider in {"demonstration", "mock", "local"}}


@router.get("/tenants/{tenant_id}/social")
def list_social(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    rows = list(db.scalars(select(CB1SocialContent).where(CB1SocialContent.tenant_id == tenant_id).order_by(CB1SocialContent.created_at.desc())))
    metas = {m.social_content_id: m for m in db.scalars(select(SocialGenerationMetadata).where(SocialGenerationMetadata.tenant_id == tenant_id))}
    return {"items": [{**_json_row(row), "generation": _json_row(metas[row.id]) if row.id in metas else None} for row in rows]}


def _campaign_recipients(db: Session, tenant_id: str, payload: CampaignExportIn) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    allowed_sources = {s.lower() for s in payload.sources}
    allowed_statuses = {s.lower() for s in payload.lead_statuses}

    def add(email: str, first_name: str, last_name: str, company: str, source: str, status: str, record_type: str) -> None:
        key = email.strip().lower()
        if not key or "@" not in key or key in seen:
            return
        if allowed_sources and source.lower() not in allowed_sources:
            return
        if allowed_statuses and status.lower() not in allowed_statuses:
            return
        seen.add(key)
        rows.append({
            "Email Address": key,
            "First Name": first_name,
            "Last Name": last_name,
            "Company": company,
            "Source": source,
            "Status": status,
            "RMR Record Type": record_type,
            "Tags": f"RMR Global,{source},{status}",
        })

    if payload.include_leads:
        for lead in db.scalars(select(Lead).where(Lead.tenant_id == tenant_id)):
            name = lead.contact_name.strip().split(maxsplit=1)
            add(lead.email, name[0] if name else "", name[1] if len(name) > 1 else "", lead.company_name, lead.source, lead.status, "Lead")
    if payload.include_contacts:
        account_names = {a.id: a.name for a in db.scalars(select(Account).where(Account.tenant_id == tenant_id))}
        for contact in db.scalars(select(Contact).where(Contact.tenant_id == tenant_id)):
            add(contact.email, contact.first_name, contact.last_name, account_names.get(contact.account_id, ""), "CRM Contact", "Active", "Contact")
    if payload.include_piq:
        for prospect in db.scalars(select(PiqOpportunity).where(PiqOpportunity.tenant_id == tenant_id)):
            # PIQ discovery records do not always include an email until a real provider enriches them.
            # Preserve them in the export only when a syntactically valid enriched email is available in the signal field.
            email_match = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", prospect.signal or "")
            if email_match:
                add(email_match.group(0), "", "", prospect.company_name, "ProspectIQ", prospect.status, "ProspectIQ")
    return rows


def _export_headers(provider: str) -> list[str]:
    if provider == "mailchimp":
        return ["Email Address", "First Name", "Last Name", "Company", "Tags", "Source", "Status", "RMR Record Type"]
    if provider == "constant_contact":
        return ["Email Address", "First Name", "Last Name", "Company", "Source", "Status", "Tags", "RMR Record Type"]
    if provider == "hubspot":
        return ["Email Address", "First Name", "Last Name", "Company", "Source", "Status", "RMR Record Type", "Tags"]
    if provider == "brevo":
        return ["Email Address", "First Name", "Last Name", "Company", "Source", "Status", "Tags", "RMR Record Type"]
    return ["Email Address", "First Name", "Last Name", "Company", "Source", "Status", "RMR Record Type", "Tags"]


@router.get("/tenants/{tenant_id}/campaign-exports")
def list_campaign_exports(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    rows = list(db.scalars(select(CampaignExportPackage).where(CampaignExportPackage.tenant_id == tenant_id).order_by(CampaignExportPackage.created_at.desc())))
    return {"items": [_json_row(row) for row in rows]}


@router.post("/tenants/{tenant_id}/campaign-exports")
def create_campaign_export(tenant_id: str, payload: CampaignExportIn, request: Request,
                           user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    require_client_campaign_write(user, tenant_id)
    provider = payload.provider_format.lower()
    if provider not in PROVIDER_EXPORTS:
        raise HTTPException(status_code=422, detail="Unsupported export provider format")
    recipients = _campaign_recipients(db, tenant_id, payload)
    export_dir = settings.data_dir / "campaign-exports" / tenant_id
    export_dir.mkdir(parents=True, exist_ok=True)
    package = CampaignExportPackage(
        tenant_id=tenant_id,
        name=payload.name,
        provider_format=provider,
        segment_json={
            "include_leads": payload.include_leads,
            "include_contacts": payload.include_contacts,
            "include_piq": payload.include_piq,
            "lead_statuses": payload.lead_statuses,
            "sources": payload.sources,
        },
        subject=payload.subject,
        preview_text=payload.preview_text,
        email_body=payload.email_body,
        call_to_action=payload.call_to_action,
        record_count=len(recipients),
        created_by=user.id,
    )
    db.add(package)
    db.flush()
    target = export_dir / f"{package.id}-{provider}.csv"
    headers = _export_headers(provider)
    with target.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(recipients)
    package.file_path = str(target)
    audit(db, user, "campaign.export.created", tenant_id=tenant_id, entity_type="campaign_export", entity_id=package.id,
          data={"provider": provider, "record_count": len(recipients), "bulk_delivery": False})
    db.commit()
    return {"package": _json_row(package), "download_url": f"/api/v521/campaign-exports/{package.id}/download"}


@router.get("/campaign-exports/{package_id}/download")
def download_campaign_export(package_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    row = db.get(CampaignExportPackage, package_id)
    if not row:
        raise HTTPException(status_code=404, detail="Campaign export not found")
    require_tenant_access(user, row.tenant_id)
    path = Path(row.file_path)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Campaign export file is unavailable")
    return FileResponse(path, filename=f"{re.sub(r'[^A-Za-z0-9_-]+', '-', row.name).strip('-')}-{row.provider_format}.csv", media_type="text/csv")


def _send_smtp(connection: CB1ProviderConnection, payload: OneToOneEmailIn) -> str:
    cfg = loads(connection.config_json, {}) or {}
    creds = loads(decrypt(connection.credential_encrypted), {}) or {}
    host = cfg.get("smtp_host") or creds.get("smtp_host")
    port = int(cfg.get("smtp_port") or creds.get("smtp_port") or 587)
    username = creds.get("username")
    password = creds.get("password")
    if not host:
        raise RuntimeError("SMTP host is not configured")
    message = EmailMessage()
    message["From"] = f"{connection.sender_name or ''} <{connection.sender_email}>" if connection.sender_name else connection.sender_email
    message["To"] = payload.recipient_email
    message["Subject"] = payload.subject
    message.set_content(payload.body)
    with smtplib.SMTP(host, port, timeout=20) as client:
        if cfg.get("starttls", True):
            client.starttls(context=ssl.create_default_context())
        if username:
            client.login(username, password or "")
        client.send_message(message)
    return message.get("Message-ID") or sha(f"{payload.recipient_email}|{payload.subject}|{_now().isoformat()}")


def _send_oauth(connection: CB1ProviderConnection, payload: OneToOneEmailIn) -> str:
    provider = connection.provider.upper()
    token = decrypt(connection.token_encrypted)
    if not token:
        raise RuntimeError("Mailbox authorization token is not configured")
    if provider in {"MICROSOFT", "OUTLOOK", "MICROSOFT365"}:
        response = httpx.post(
            "https://graph.microsoft.com/v1.0/me/sendMail",
            headers={"Authorization": f"Bearer {token}"},
            json={"message": {"subject": payload.subject, "body": {"contentType": "Text", "content": payload.body}, "toRecipients": [{"emailAddress": {"address": payload.recipient_email}}]}, "saveToSentItems": True},
            timeout=30,
        )
        response.raise_for_status()
        return response.headers.get("request-id", sha(f"ms|{payload.recipient_email}|{_now().isoformat()}"))
    if provider in {"GOOGLE", "GMAIL", "GOOGLE_WORKSPACE"}:
        message = EmailMessage()
        message["From"] = connection.sender_email
        message["To"] = payload.recipient_email
        message["Subject"] = payload.subject
        message.set_content(payload.body)
        raw = base64.urlsafe_b64encode(message.as_bytes()).decode().rstrip("=")
        response = httpx.post(
            "https://gmail.googleapis.com/gmail/v1/users/me/messages/send",
            headers={"Authorization": f"Bearer {token}"},
            json={"raw": raw},
            timeout=30,
        )
        response.raise_for_status()
        return str(response.json().get("id") or sha(f"google|{payload.recipient_email}|{_now().isoformat()}"))
    raise RuntimeError(f"Provider {provider} does not support direct one-to-one sending")


@router.get("/tenants/{tenant_id}/email-workspace")
def email_workspace(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    connections = list(db.scalars(select(CB1ProviderConnection).where(CB1ProviderConnection.tenant_id == tenant_id).order_by(CB1ProviderConnection.created_at.desc())))
    messages = list(db.scalars(select(CB1Message).where(CB1Message.tenant_id == tenant_id).order_by(CB1Message.created_at.desc()).limit(100)))
    activities = list(db.scalars(select(Activity).where(Activity.tenant_id == tenant_id).order_by(Activity.created_at.desc()).limit(100)))
    contexts = {row.message_id: row for row in db.scalars(select(MessageDeliveryContext).where(MessageDeliveryContext.tenant_id == tenant_id))}
    user_ids = {row.created_by for row in messages if row.created_by}
    users = {row.id: row for row in db.scalars(select(User).where(User.id.in_(user_ids)))} if user_ids else {}
    safe_connections = [_json_row(row, ("token_encrypted", "refresh_token_encrypted", "credential_encrypted")) for row in connections]
    connection_map = {row.id: row for row in connections}
    message_rows = []
    for row in messages:
        context = contexts.get(row.id)
        connection = connection_map.get(context.connection_id) if context else None
        actor = users.get(row.created_by)
        message_rows.append(_json_row(row, ("body_draft", "body_approved", "supported_facts_json")) | {
            "body": row.body_approved or row.body_draft,
            "sender_email": context.sender_email if context else (connection.sender_email if connection else ""),
            "sender_name": context.sender_name if context else (connection.sender_name if connection else ""),
            "provider": context.provider if context else (connection.provider if connection else "unknown"),
            "recipient_name": context.recipient_name if context else "",
            "connection_id": context.connection_id if context else None,
            "delivery_status": context.delivery_status if context else row.status,
            "actor_name": actor.full_name if actor else "",
            "actor_email": actor.email if actor else "",
            "related_account_id": context.account_id if context else None,
            "related_opportunity_id": context.opportunity_id if context else None,
            "related_lead_id": context.lead_id if context else None,
            "related_contact_id": context.contact_id if context else None,
        })
    return {
        "connections": safe_connections,
        "messages": message_rows,
        "activities": [_json_row(row) for row in activities],
        "policy": {
            "one_to_one_only": True,
            "client_owned_sender": True,
            "bulk_delivery": False,
            "detail": "One-to-one CRM email is sent from the client's connected mailbox. Bulk campaigns are exported to a dedicated email platform.",
        },
    }


@router.post("/tenants/{tenant_id}/email-connections/{connection_id}/action")
def email_connection_action(tenant_id: str, connection_id: str, payload: dict[str, Any], request: Request,
                            user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    _client_admin_or_managed(user, tenant_id)
    row = db.get(CB1ProviderConnection, connection_id)
    if not row or row.tenant_id != tenant_id:
        raise HTTPException(status_code=404, detail="Email connection not found")
    action = str(payload.get("action", "test")).lower()
    detail = ""
    if action == "disconnect":
        row.status = "DISCONNECTED"
        row.token_encrypted = None
        row.refresh_token_encrypted = None
        row.credential_encrypted = None
        detail = "Mailbox connection disconnected."
    elif action == "sync":
        row.last_sync_at = _now()
        detail = "CRM-relevant mailbox synchronization recorded. Production providers must supply message sync permissions."
    elif action == "test":
        provider = row.provider.upper()
        if provider == "MOCK":
            row.status = "ACTIVE"
            detail = "Demonstration mailbox is operational."
        elif provider in {"SMTP", "IMAP", "SMTP_IMAP"}:
            cfg = loads(row.config_json, {}) or {}
            creds = loads(decrypt(row.credential_encrypted), {}) or {}
            host = cfg.get("smtp_host") or creds.get("smtp_host")
            if not host:
                raise HTTPException(status_code=409, detail="SMTP host is not configured")
            row.status = "SAVED_UNTESTED"
            detail = "SMTP configuration is present. Use production connectivity to complete authentication testing."
        else:
            row.status = "ACTIVE" if decrypt(row.token_encrypted) else "AUTHORIZATION_REQUIRED"
            detail = "OAuth authorization is present." if row.status == "ACTIVE" else "Provider authorization is required."
    else:
        raise HTTPException(status_code=422, detail="Unsupported email connection action")
    event = EmailConnectionEvent(tenant_id=tenant_id, connection_id=connection_id, event_type=action, status=row.status, detail=detail, created_by=user.id)
    db.add(event)
    audit(db, user, f"email.connection.{action}", tenant_id=tenant_id, entity_type="provider_connection", entity_id=row.id, data={"status": row.status})
    db.commit()
    return {"connection": _json_row(row, ("token_encrypted", "refresh_token_encrypted", "credential_encrypted")), "detail": detail}


@router.post("/tenants/{tenant_id}/one-to-one-email")
def send_one_to_one_email(tenant_id: str, payload: OneToOneEmailIn, request: Request,
                          user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    require_client_operational_write(user, tenant_id)
    if "@" not in payload.recipient_email:
        raise HTTPException(status_code=422, detail="Enter a valid recipient email")
    connection = db.get(CB1ProviderConnection, payload.connection_id)
    if not connection or connection.tenant_id != tenant_id:
        raise HTTPException(status_code=404, detail="Connected mailbox not found")
    provider = connection.provider.upper()
    provider_message_id = ""
    delivery_status = "RECORDED_DEMONSTRATION"
    try:
        if provider == "MOCK":
            provider_message_id = sha(f"mock|{payload.recipient_email}|{payload.subject}|{_now().isoformat()}")
        elif provider in {"SMTP", "IMAP", "SMTP_IMAP"}:
            provider_message_id = _send_smtp(connection, payload)
            delivery_status = "SENT"
        else:
            provider_message_id = _send_oauth(connection, payload)
            delivery_status = "SENT"
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Client mailbox send failed: {exc}")
    key = sha("|".join([tenant_id, payload.connection_id, payload.recipient_email.lower(), payload.subject, payload.body, _now().isoformat()]))
    message = CB1Message(
        tenant_id=tenant_id,
        crm_record_id=payload.opportunity_id or payload.account_id or payload.lead_id or payload.contact_id,
        recipient_email=payload.recipient_email.lower(),
        subject=payload.subject,
        body_draft=payload.body,
        body_approved=payload.body,
        approved_hash=sha(payload.body),
        status=delivery_status,
        approved_by=user.id,
        approved_at=_now(),
        sent_at=_now(),
        provider_message_id=provider_message_id,
        idempotency_key=key,
        created_by=user.id,
    )
    db.add(message)
    db.flush()
    db.add(MessageDeliveryContext(
        message_id=message.id,
        tenant_id=tenant_id,
        connection_id=connection.id,
        provider=provider,
        sender_email=connection.sender_email,
        sender_name=connection.sender_name or "",
        recipient_name=payload.recipient_name,
        account_id=payload.account_id,
        opportunity_id=payload.opportunity_id,
        lead_id=payload.lead_id,
        contact_id=payload.contact_id,
        delivery_status=delivery_status,
        created_by=user.id,
    ))
    activity = Activity(
        tenant_id=tenant_id,
        account_id=payload.account_id,
        opportunity_id=payload.opportunity_id,
        user_id=user.id,
        activity_type="Email",
        subject=payload.subject,
        body=f"To: {payload.recipient_name or payload.recipient_email} <{payload.recipient_email}>\n\n{payload.body}",
        completed_at=_now(),
    )
    db.add(activity)
    audit(db, user, "crm.one_to_one_email.sent", tenant_id=tenant_id, entity_type="message", entity_id=message.id,
          data={"provider": provider, "status": delivery_status, "account_id": payload.account_id, "opportunity_id": payload.opportunity_id})
    db.commit()
    return {"message": _json_row(message, ("supported_facts_json",)), "activity": _json_row(activity), "delivery_status": delivery_status}


def _parse_money(value: Any) -> int:
    text_value = str(value or "").strip()
    if not text_value:
        return 0
    negative = text_value.startswith("(") and text_value.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text_value)
    amount = float(cleaned or 0)
    if negative:
        amount = -abs(amount)
    return int(round(amount * 100))


def _parse_month(value: Any) -> int:
    text_value = str(value or "").strip()
    if not text_value:
        raise ValueError("Month is required")
    month_map = {name.lower(): index for index, name in enumerate(["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"], 1)}
    month_map.update({name[:3].lower(): index for name, index in month_map.items()})
    if text_value.lower() in month_map:
        return month_map[text_value.lower()]
    number_value = int(float(text_value))
    if not 1 <= number_value <= 12:
        raise ValueError("Month must be 1-12 or a month name")
    return number_value


def _normalize_target(value: Any) -> str:
    text_value = str(value or "prior_actual").strip().lower().replace(" ", "_").replace("-", "_")
    aliases = {
        "prior": "prior_actual",
        "prior_year": "prior_actual",
        "prior_actual": "prior_actual",
        "actual": "actual",
        "actual_to_date": "actual",
        "forecast": "forecast",
        "operating_forecast": "forecast",
    }
    if text_value not in aliases:
        raise ValueError("Target must be prior_actual, actual, or forecast")
    return aliases[text_value]


def _read_forecast_file(upload: UploadFile) -> list[dict[str, Any]]:
    raw = upload.file.read()
    suffix = Path(upload.filename or "").suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="XLSX support is unavailable") from exc
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(v or "").strip() for v in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(v not in {None, ""} for v in row)]
    text_value = raw.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text_value)))


def _normalize_forecast_rows(db: Session, tenant_id: str, rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    accounts = {a.name.strip().lower(): a for a in db.scalars(select(Account).where(Account.tenant_id == tenant_id))}
    normalized: list[dict[str, Any]] = []
    exceptions: list[dict[str, Any]] = []
    for index, row in enumerate(rows, 2):
        lower = {str(k).strip().lower().replace(" ", "_"): v for k, v in row.items()}
        try:
            account_name = str(lower.get("account") or lower.get("account_name") or "").strip()
            account = accounts.get(account_name.lower())
            if not account:
                raise ValueError(f"Account not matched: {account_name or '(blank)'}")
            month_number = _parse_month(lower.get("month"))
            amount_cents = _parse_money(lower.get("amount") or lower.get("value"))
            target = _normalize_target(lower.get("target") or lower.get("type"))
            normalized.append({
                "row": index,
                "account_id": account.id,
                "account": account.name,
                "month": month_number,
                "amount_cents": amount_cents,
                "target": target,
                "display_amount": f"${amount_cents / 100:,.2f}",
            })
        except Exception as exc:
            exceptions.append({"row": index, "reason": str(exc), "source": lower})
    return normalized, exceptions


@router.post("/tenants/{tenant_id}/forecast-import/preview")
def preview_forecast_import(tenant_id: str, version_id: str = Form(...), file: UploadFile = File(...),
                            user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_client_operational_write(user, tenant_id)
    version = db.get(ForecastVersion, version_id)
    if not version or version.tenant_id != tenant_id:
        raise HTTPException(status_code=400, detail="Invalid forecast version")
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".csv", ".xlsx", ".xlsm"}:
        raise HTTPException(status_code=422, detail="Upload a CSV or XLSX file")
    rows, exceptions = _normalize_forecast_rows(db, tenant_id, _read_forecast_file(file))
    batch = ForecastImportBatch(tenant_id=tenant_id, version_id=version_id, file_name=file.filename or "forecast-import", rows_json=rows, exceptions_json=exceptions, created_by=user.id)
    db.add(batch)
    db.commit()
    return {"batch": _json_row(batch), "rows": rows, "exceptions": exceptions, "ready_to_import": bool(rows)}


@router.post("/tenants/{tenant_id}/forecast-import/commit")
def commit_forecast_import(tenant_id: str, payload: ForecastCommitIn, request: Request,
                           user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    require_client_operational_write(user, tenant_id)
    batch = db.get(ForecastImportBatch, payload.batch_id)
    if not batch or batch.tenant_id != tenant_id or batch.status != "preview":
        raise HTTPException(status_code=404, detail="Forecast import preview is unavailable")
    changed: list[dict[str, Any]] = []
    for row in batch.rows_json:
        forecast = db.scalar(select(ForecastMonth).where(
            ForecastMonth.version_id == batch.version_id,
            ForecastMonth.account_id == row["account_id"],
            ForecastMonth.month == int(row["month"]),
        ))
        if not forecast:
            forecast = ForecastMonth(version_id=batch.version_id, account_id=row["account_id"], month=int(row["month"]))
            db.add(forecast)
            db.flush()
        target = row["target"]
        field = {"prior_actual": "prior_actual_cents", "actual": "actual_cents", "forecast": "forecast_cents"}[target]
        before = int(getattr(forecast, field) or 0)
        setattr(forecast, field, int(row["amount_cents"]))
        changed.append({**row, "before_cents": before, "after_cents": int(row["amount_cents"])})
    batch.status = "committed"
    batch.committed_at = _now()
    audit(db, user, "forecast.file_import.committed", tenant_id=tenant_id, entity_type="forecast_import", entity_id=batch.id,
          data={"changed_rows": len(changed), "exception_rows": len(batch.exceptions_json)})
    db.commit()
    return {"batch": _json_row(batch), "changed_rows": changed, "exceptions": batch.exceptions_json}


@router.get("/tenants/{tenant_id}/forecast-import/template")
def forecast_template(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    accounts = list(db.scalars(select(Account).where(Account.tenant_id == tenant_id).order_by(Account.name)))
    content = io.StringIO()
    writer = csv.writer(content)
    writer.writerow(["Account", "Month", "Amount", "Target"])
    for account in accounts[:3]:
        writer.writerow([account.name, "January", "25000.00", "prior_actual"])
    path = settings.data_dir / "forecast-import-template.csv"
    path.write_text(content.getvalue(), encoding="utf-8-sig")
    return FileResponse(path, filename="RMR-Global-Forecast-Import-Template.csv", media_type="text/csv")


@router.get("/tenants/{tenant_id}/forecast-provenance")
def forecast_provenance(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    version = db.scalar(select(ForecastVersion).where(ForecastVersion.tenant_id == tenant_id, ForecastVersion.is_active.is_(True)).order_by(ForecastVersion.created_at.desc()))
    closed_won = db.scalar(select(func.coalesce(func.sum(Opportunity.value_cents), 0)).where(Opportunity.tenant_id == tenant_id, func.lower(Opportunity.stage).in_(["closed won", "won"]))) or 0
    imported_actuals = 0
    if version:
        imported_actuals = db.scalar(select(func.coalesce(func.sum(ForecastMonth.actual_cents), 0)).where(ForecastMonth.version_id == version.id)) or 0
    return {
        "actual_to_date_cents": int(imported_actuals),
        "actual_source": "Imported actual sales in the active forecast version",
        "closed_won_revenue_cents": int(closed_won),
        "closed_won_source": "CRM opportunities whose stage is Closed Won",
        "forecast_source": "Client-entered operating forecast and committed forecast imports",
        "annual_goal_source": "Client-entered annual goal on the active forecast version",
    }


@router.get("/tenants/{tenant_id}/action-report")
def action_report(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    website_leads = db.scalar(select(func.count()).select_from(Lead).where(Lead.tenant_id == tenant_id, Lead.source == "Website")) or 0
    piq_records = db.scalar(select(func.count()).select_from(PiqOpportunity).where(PiqOpportunity.tenant_id == tenant_id)) or 0
    piq_to_crm = db.scalar(select(func.count()).select_from(PiqOpportunity).where(PiqOpportunity.tenant_id == tenant_id, PiqOpportunity.moved_to_crm.is_(True))) or 0
    open_opps = db.scalar(select(func.count()).select_from(Opportunity).where(Opportunity.tenant_id == tenant_id, ~func.lower(Opportunity.stage).in_(["closed won", "won", "closed lost", "lost"]))) or 0
    won_opps = db.scalar(select(func.count()).select_from(Opportunity).where(Opportunity.tenant_id == tenant_id, func.lower(Opportunity.stage).in_(["closed won", "won"]))) or 0
    won_revenue = db.scalar(select(func.coalesce(func.sum(Opportunity.value_cents), 0)).where(Opportunity.tenant_id == tenant_id, func.lower(Opportunity.stage).in_(["closed won", "won"]))) or 0
    social_drafts = db.scalar(select(func.count()).select_from(CB1SocialContent).where(CB1SocialContent.tenant_id == tenant_id, CB1SocialContent.status == "DRAFT")) or 0
    messages = db.scalar(select(func.count()).select_from(CB1Message).where(CB1Message.tenant_id == tenant_id)) or 0
    attention = []
    if piq_records > piq_to_crm:
        attention.append({"title": f"{piq_records - piq_to_crm} ProspectIQ records need review", "detail": "Review qualified prospects and decide which should enter CRM.", "route": "piq", "label": "Review ProspectIQ"})
    if open_opps:
        attention.append({"title": f"{open_opps} open opportunities", "detail": "Review next actions, expected close dates, and stale activity.", "route": "crm?tab=opportunities", "label": "Review Pipeline"})
    if social_drafts:
        attention.append({"title": f"{social_drafts} social drafts awaiting review", "detail": "Edit, approve, copy, or record publication.", "route": "campaigns?tab=social", "label": "Review Drafts"})
    if website_leads:
        attention.append({"title": f"{website_leads} website leads", "detail": "Confirm follow-up and conversion status.", "route": "crm?tab=leads&source=Website", "label": "Open Website Leads"})
    source_rows = list(db.execute(select(Opportunity.source, func.count(Opportunity.id), func.coalesce(func.sum(Opportunity.value_cents), 0)).where(Opportunity.tenant_id == tenant_id).group_by(Opportunity.source)))
    return {
        "kpis": [
            {"key": "website_leads", "label": "Website Leads", "value": int(website_leads), "route": "crm?tab=leads&source=Website"},
            {"key": "piq_prospects", "label": "PIQ Prospects", "value": int(piq_records), "route": "piq"},
            {"key": "piq_to_crm", "label": "PIQ to CRM", "value": int(piq_to_crm), "route": "crm?tab=leads&source=ProspectIQ"},
            {"key": "open_opportunities", "label": "Open Opportunities", "value": int(open_opps), "route": "crm?tab=opportunities"},
            {"key": "won_opportunities", "label": "Won Opportunities", "value": int(won_opps), "route": "crm?tab=opportunities&stage=Won"},
            {"key": "won_revenue", "label": "Won Revenue", "value_cents": int(won_revenue), "route": "crm?tab=opportunities&stage=Won"},
            {"key": "social_drafts", "label": "Social Drafts", "value": int(social_drafts), "route": "campaigns?tab=social"},
            {"key": "emails", "label": "One-to-One Emails", "value": int(messages), "route": "email?tab=messages"},
        ],
        "attention": attention,
        "funnel": [
            {"label": "Website Leads", "value": int(website_leads), "route": "crm?tab=leads&source=Website"},
            {"label": "PIQ Records", "value": int(piq_records), "route": "piq"},
            {"label": "PIQ Moved to CRM", "value": int(piq_to_crm), "route": "crm?tab=leads&source=ProspectIQ"},
            {"label": "Open Opportunities", "value": int(open_opps), "route": "crm?tab=opportunities"},
            {"label": "Won Opportunities", "value": int(won_opps), "route": "crm?tab=opportunities&stage=Won"},
        ],
        "sources": [{"source": source or "Unknown", "opportunities": int(count), "value_cents": int(value), "route": f"crm?tab=opportunities&source={source or 'Unknown'}"} for source, count, value in source_rows],
    }


@router.get("/tenants/{tenant_id}/client-training")
def list_client_training(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    resources = list(db.scalars(select(ClientTrainingResource).where(
        or_(ClientTrainingResource.tenant_id == tenant_id, ClientTrainingResource.owner_scope == "rmr"),
        ClientTrainingResource.archived.is_(False),
    ).order_by(ClientTrainingResource.created_at.desc())))
    assignments = list(db.scalars(select(ClientTrainingAssignment).where(ClientTrainingAssignment.tenant_id == tenant_id)))
    users = list(db.scalars(select(User).where(User.tenant_id == tenant_id, User.active.is_(True)).order_by(User.full_name)))
    by_resource: dict[str, list[dict[str, Any]]] = {}
    user_map = {row.id: row for row in users}
    for assignment in assignments:
        target = user_map.get(assignment.user_id)
        by_resource.setdefault(assignment.resource_id, []).append({**_json_row(assignment), "user_name": target.full_name if target else "Unknown user"})
    return {
        "resources": [{**_json_row(row), "assignments": by_resource.get(row.id, [])} for row in resources],
        "users": [{"id": row.id, "name": row.full_name, "email": row.email, "role": row.tenant_role} for row in users],
        "policy": "Clients may upload or link workflow training for their staff. RMR platform training can be added later under the separate RMR scope.",
    }


@router.post("/tenants/{tenant_id}/client-training")
def create_client_training(
    tenant_id: str,
    title: str = Form(...),
    description: str = Form(""),
    category: str = Form("Sales Workflow"),
    media_url: str = Form(""),
    required: bool = Form(False),
    roles_json: str = Form("[]"),
    user_ids_json: str = Form("[]"),
    due_date: str = Form(""),
    file: UploadFile | None = File(None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    _client_admin_or_managed(user, tenant_id)
    roles = json.loads(roles_json or "[]")
    user_ids = json.loads(user_ids_json or "[]")
    file_path = ""
    media_type = "external_link"
    if file and file.filename:
        suffix = Path(file.filename).suffix.lower()
        allowed = {".mp4", ".webm", ".mov", ".m4v", ".pdf", ".pptx", ".docx"}
        if suffix not in allowed:
            raise HTTPException(status_code=422, detail="Upload a supported video, PDF, PowerPoint, or Word training file")
        target_dir = settings.data_dir / "client-training" / tenant_id
        target_dir.mkdir(parents=True, exist_ok=True)
        safe_name = f"{datetime.now().strftime('%Y%m%d%H%M%S')}-{re.sub(r'[^A-Za-z0-9_.-]', '-', Path(file.filename).name)}"
        target = target_dir / safe_name
        with target.open("wb") as handle:
            while chunk := file.file.read(1024 * 1024):
                handle.write(chunk)
        file_path = str(target)
        media_type = "uploaded_video" if suffix in {".mp4", ".webm", ".mov", ".m4v"} else "uploaded_file"
    elif media_url.strip():
        _safe_url(media_url)
    else:
        raise HTTPException(status_code=422, detail="Upload a training file or provide a hosted training URL")
    parsed_due = date.fromisoformat(due_date) if due_date else None
    resource = ClientTrainingResource(
        tenant_id=tenant_id,
        owner_scope="client",
        title=title.strip(),
        description=description.strip(),
        category=category.strip() or "Sales Workflow",
        media_type=media_type,
        media_url=media_url.strip(),
        file_path=file_path,
        required=required,
        roles_json=[str(x) for x in roles],
        due_date=parsed_due,
        created_by=user.id,
    )
    db.add(resource)
    db.flush()
    target_users = list(db.scalars(select(User).where(User.tenant_id == tenant_id, User.active.is_(True))))
    for target_user in target_users:
        if user_ids and target_user.id not in user_ids:
            continue
        if roles and target_user.tenant_role not in roles:
            continue
        db.add(ClientTrainingAssignment(tenant_id=tenant_id, resource_id=resource.id, user_id=target_user.id))
    audit(db, user, "client.training.created", tenant_id=tenant_id, entity_type="client_training", entity_id=resource.id, data={"category": resource.category, "required": required})
    db.commit()
    return {"resource": _json_row(resource)}


@router.get("/client-training/{resource_id}/media")
def client_training_media(resource_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    resource = db.get(ClientTrainingResource, resource_id)
    if not resource or not resource.file_path:
        raise HTTPException(status_code=404, detail="Training file not found")
    require_tenant_access(user, resource.tenant_id or user.tenant_id or "")
    path = Path(resource.file_path)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Training file is unavailable")
    return FileResponse(path, media_type=mimetypes.guess_type(path.name)[0] or "application/octet-stream", filename=path.name)


@router.patch("/client-training/assignments/{assignment_id}")
def update_training_assignment(assignment_id: str, payload: TrainingCompleteIn, request: Request,
                               user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    row = db.get(ClientTrainingAssignment, assignment_id)
    if not row:
        raise HTTPException(status_code=404, detail="Training assignment not found")
    if not is_global_admin(user) and row.user_id != user.id and user.tenant_role != "CLIENT_ADMIN":
        raise HTTPException(status_code=403, detail="Training assignment access denied")
    require_tenant_access(user, row.tenant_id)
    row.status = payload.status
    row.progress_pct = payload.progress_pct
    row.completed_at = _now() if payload.status == "complete" else None
    audit(db, user, "client.training.progress", tenant_id=row.tenant_id, entity_type="training_assignment", entity_id=row.id, data={"status": row.status, "progress_pct": row.progress_pct})
    db.commit()
    return {"assignment": _json_row(row)}


@router.delete("/client-training/{resource_id}")
def archive_client_training(resource_id: str, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    row = db.get(ClientTrainingResource, resource_id)
    if not row or not row.tenant_id:
        raise HTTPException(status_code=404, detail="Client training resource not found")
    _client_admin_or_managed(user, row.tenant_id)
    row.archived = True
    audit(db, user, "client.training.archived", tenant_id=row.tenant_id, entity_type="client_training", entity_id=row.id)
    db.commit()
    return {"ok": True}


def _send_system_email(recipient: str, subject: str, body: str) -> str:
    provider = os.getenv("RMR_SYSTEM_EMAIL_PROVIDER", "local").lower()
    if provider == "smtp":
        host = os.getenv("RMR_SYSTEM_SMTP_HOST")
        if not host:
            raise RuntimeError("RMR_SYSTEM_SMTP_HOST is not configured")
        port = int(os.getenv("RMR_SYSTEM_SMTP_PORT", "587"))
        sender = os.getenv("RMR_SYSTEM_EMAIL_FROM", "no-reply@rmr.invalid")
        message = EmailMessage()
        message["From"] = sender
        message["To"] = recipient
        message["Subject"] = subject
        message.set_content(body)
        with smtplib.SMTP(host, port, timeout=20) as client:
            if os.getenv("RMR_SYSTEM_SMTP_STARTTLS", "true").lower() == "true":
                client.starttls(context=ssl.create_default_context())
            username = os.getenv("RMR_SYSTEM_SMTP_USERNAME")
            if username:
                client.login(username, os.getenv("RMR_SYSTEM_SMTP_PASSWORD", ""))
            client.send_message(message)
        return "sent"
    if settings.environment == "production":
        raise RuntimeError("System transactional email provider is not configured")
    return "local_recovery"


@router.post("/tenants/{tenant_id}/users/send-password-reset")
def send_user_password_reset(tenant_id: str, payload: PasswordResetSendIn, request: Request,
                             user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    _client_admin_or_managed(user, tenant_id)
    target = db.get(User, payload.user_id)
    if not target or target.tenant_id != tenant_id:
        raise HTTPException(status_code=404, detail="Team member not found")
    reset, raw_token = create_password_reset(db, target, request.client.host if request.client else "")
    url = password_reset_url(raw_token)
    delivery = _send_system_email(target.email, "Reset your RMR Global password", f"A password reset was requested for your RMR Global account.\n\n{url}\n\nThis single-use link expires in {settings.password_reset_minutes} minutes.")
    local_url = None
    if delivery == "local_recovery":
        write_local_recovery_file(raw_token, target.email)
        local_url = url
    audit(db, user, "team.password_reset.sent", tenant_id=tenant_id, entity_type="user", entity_id=target.id, data={"delivery": delivery, "expires_at": reset.expires_at.isoformat()})
    db.commit()
    return {"sent": delivery == "sent", "delivery": delivery, "email": target.email, "expires_at": reset.expires_at.isoformat(), "local_reset_url": local_url}


@router.post("/tenants/{tenant_id}/solution-requests")
def corrected_solution_request(tenant_id: str, payload: SolutionRequestIn, request: Request,
                               user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_request_origin(request)
    _client_admin_or_managed(user, tenant_id)
    catalog = db.scalar(select(ServiceCatalog).where(ServiceCatalog.code == payload.service_code, ServiceCatalog.active.is_(True)))
    if not catalog:
        raise HTTPException(status_code=404, detail="Solution not found")
    active = db.scalar(select(TenantService).where(TenantService.tenant_id == tenant_id, TenantService.service_code == payload.service_code, TenantService.status == "active"))
    if active:
        raise HTTPException(status_code=409, detail="This solution is already active")
    existing = db.scalar(select(SolutionRequest).where(
        SolutionRequest.tenant_id == tenant_id,
        SolutionRequest.service_code == payload.service_code,
        SolutionRequest.status.in_(["Requested", "Pending Activation"]),
    ))
    if existing:
        preference = db.get(SolutionRequestPreference, existing.id)
        return {"request": _json_row(existing), "preference": _json_row(preference) if preference else None, "duplicate_prevented": True}
    request_row = SolutionRequest(
        tenant_id=tenant_id,
        service_code=payload.service_code,
        requested_by=user.id,
        status="Requested",
        note=payload.note,
        proposed_monthly_cents=catalog.standard_price_cents if catalog.cadence == "monthly" else 0,
        proposed_usage_cents=catalog.standard_price_cents if catalog.cadence == "usage" else 0,
        preferred_contact_method=payload.preferred_contact_method,
        best_time=f"{payload.contact_date.isoformat()} {payload.contact_time} {payload.timezone}",
    )
    db.add(request_row)
    db.flush()
    preference = SolutionRequestPreference(
        request_id=request_row.id,
        tenant_id=tenant_id,
        contact_date=payload.contact_date,
        contact_time=payload.contact_time,
        timezone=payload.timezone,
        confirmed=False,
    )
    db.add(preference)
    db.add(Notification(
        recipient_scope="GLOBAL_ADMIN",
        tenant_id=tenant_id,
        notification_type="solution_request",
        title=f"New solution request — {catalog.name}",
        body=f"Client requested {catalog.name}; preferred contact: {payload.contact_date.isoformat()} {payload.contact_time} {payload.timezone}. This is not a confirmed appointment.",
        action_route="service-requests",
        action_label="Review request",
        entity_type="solution_request",
        entity_id=request_row.id,
    ))
    audit(db, user, "solution.requested.with_contact_preference", tenant_id=tenant_id, entity_type="solution_request", entity_id=request_row.id,
          data={"service_code": payload.service_code, "contact_date": payload.contact_date.isoformat(), "contact_time": payload.contact_time, "timezone": payload.timezone, "automatic_activation": False})
    db.commit()
    return {"request": _json_row(request_row), "preference": _json_row(preference), "duplicate_prevented": False,
            "confirmation": "Request received. RMR will contact you to confirm; no service was activated and no charge was created."}


@router.get("/tenants/{tenant_id}/solution-request-status")
def solution_request_status(tenant_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_tenant_access(user, tenant_id)
    rows = list(db.scalars(select(SolutionRequest).where(SolutionRequest.tenant_id == tenant_id).order_by(SolutionRequest.created_at.desc())))
    preferences = {row.request_id: row for row in db.scalars(select(SolutionRequestPreference).where(SolutionRequestPreference.tenant_id == tenant_id))}
    return {"items": [{**_json_row(row), "preference": _json_row(preferences[row.id]) if row.id in preferences else None} for row in rows]}
